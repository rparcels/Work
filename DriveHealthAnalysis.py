import tkinter as tk
from tkinter import ttk, messagebox
import json
from ZDT_SDK import ZdtClient
import pandas as pd
from pathlib import Path
from openpyxl.styles import PatternFill
from openpyxl import load_workbook
from datetime import datetime
import os


# ===============================================
# Report name/destination 
# ===============================================
REPORT_NAME = "Drive Health Analysis.xlsx"
dest_dir = Path.home() / "Documents" / "API Reports" / "Drive Health Analysis"
dest_dir.mkdir(parents=True, exist_ok=True)

OUTPUT_XLSX = dest_dir / REPORT_NAME


# ===============================================
# API Key Repository (saved in windows environmental variables)
# ===============================================
#Request variable values
DS_KEY = os.getenv("DS_API_KEY")
DWSA_KEY = os.getenv("DWSA_API_KEY")
MW_KEY = os.getenv("MW_API_KEY")
RP_KEY = os.getenv("RP_API_KEY")

#Store in dictionary
users = {
    "Danielle Sutton"           :  DS_KEY,
    "DemoWorks Service Account" :  DWSA_KEY,
    "Missy Walker"              :  MW_KEY,
    "Ricky Parcels"             :  RP_KEY
    }


# ===============================================
# Open and read config.json to see who is currently "logged in"
# ===============================================
def load_last_user():
    try:
        with open("config.json", "r") as f:
            return json.load(f).get("last_user")
    except FileNotFoundError:
        return None


def save_last_user(user):
    with open("config.json", "w") as f:
        json.dump({"last_user": user}, f)


def get_api_key(user):
    return users.get(user)


current_user = load_last_user()
current_api_key = get_api_key(current_user)


# ===============================================
# FUNCTION TO QUERY PLACES ENDPOINT (COMPANIES)
# ===============================================
def load_locations():
    with ZdtClient(api_key=current_api_key) as client:
        places = client.places.get_places()

    return {
        p.place_name: str(p.place_guid)  #Create key-value pairs. Left = Key, Right = Value
        for p in places.root             #Iterate through every place
        if p.place_type == 2             #Only include places where the type equals 2
    }



# ===============================================
# FUNCTION TO REFRESH PLACES (COMPANIES) WHEN A DIFFERENT USER IS SELECTED
# ===============================================
locations = {}

def refresh_locations():
    global locations

    if not current_api_key:
        return

    locations = load_locations()

    company_combo["values"] = sorted(locations.keys(), key=str.lower)

    if locations:
        company_combo.current(0)



# ===============================================
# ROBOT INPUTS
# ===============================================
def run_report():
    try:
        placeGuid = locations[company_var.get()]
        serial_input = serial_entry.get().strip()

        
        startDate = start_entry.get().strip()
        endDate = end_entry.get().strip()

        # =========================
        # JOB FILTER FUNCTION
        # =========================
        def robot_job_filter(r):
            return (r.serial_number == serial_input) 


        # =========================
        # QUERY CYCLE STATUS ENDPOINT w/ JOB FILTER
        # =========================
        with ZdtClient(api_key=current_api_key) as client:
            cycles = client.cycle_status.get_cycle_status_by_place(
                place_type=2,
                place_guid=placeGuid,
                start_date=startDate,
                end_date=endDate,
                company_guid=None,
                predicate=robot_job_filter  # Apply robot filter
            )


        # =========================
        # OUTPUT CHECK
        # =========================
        all_cycles_df = cycles.to_pandas()
        print(f"\nRecords returned: {len(all_cycles_df)}")

        #If no records are found, return (exit the run_report() function)
        if len(all_cycles_df) == 0:
            print("No data found for given inputs.")
            messagebox.showwarning(
                "No Data",
                "No records were found for the selected inputs."
            )
            return


        # =========================
        # Convert to dataframe 
        # =========================
        all_cycles_df = all_cycles_df.drop(columns=["ip_address", "flex_data_items", "flex_io_data_items", "power_consumption"])
        all_cycles_df["date_time"] = (pd.to_datetime(all_cycles_df["date_time"], format="mixed", utc=True, errors="coerce").dt.tz_localize(None).dt.strftime("%m/%d/%Y %I:%M").str.lstrip("0"))


        def analyze_job(job_df):
            min_ct = job_df["cycle_time"].min()
            median_ct = job_df["cycle_time"].median()
            mean_ct = job_df["cycle_time"].mean()
            max_ct = job_df["cycle_time"].max()
            std_ct = job_df["cycle_time"].std()

            print(f"\nRecords: {len(job_df)}")
            print(f"Min:    {round(min_ct,1)}")
            print(f"Median: {round(median_ct,1)}")
            print(f"Max:    {round(max_ct,1)}")
            print(f"STD:    {round(std_ct,1)}")

            # =========================
            # Z-Score
            # =========================
            if pd.isna(std_ct) or std_ct == 0:
                zscore = pd.Series(0, index=job_df.index)
            else:
                zscore = (job_df["cycle_time"] - mean_ct) / std_ct

            # =========================
            # IQR
            # =========================
            q1 = job_df["cycle_time"].quantile(0.25)
            q3 = job_df["cycle_time"].quantile(0.75)

            iqr = q3 - q1

            iqr_outlier = (
                (job_df["cycle_time"] < q1 - 1.5 * iqr)
                |
                (job_df["cycle_time"] > q3 + 1.5 * iqr)
            )

            # =========================
            # Outlier flag
            # =========================
            job_df["Outlier"] = (
                (zscore.abs() > 3)
                |
                (iqr_outlier)
            )

            # =========================
            # Range ratio
            # =========================
            p5 = job_df["cycle_time"].quantile(0.05)
            p95 = job_df["cycle_time"].quantile(0.95)

            range_ratio = (
                p95 / p5
                if p5 > 0
                else float("inf")
            )

            outlier_count = int(job_df["Outlier"].sum())
            total_count = len(job_df)

            outlier_ratio = (
                outlier_count / total_count
                if total_count > 0
                else 0
            )

            # =========================
            # Status
            # =========================
            if range_ratio >= 20:
                status = "Misconfigured Trigger"

            elif range_ratio >= 8:
                status = "Suspicious / Needs Review"

            else:

                if outlier_count == 0:
                    status = "Stable"

                elif outlier_ratio < 0.02:
                    status = f"Stable ({outlier_count} minor outliers)"

                else:
                    status = f"Review ({outlier_count} outliers)"

            return {
                "df": job_df,
                "status": status,
                "count": total_count,
                "min": min_ct,
                "median": median_ct,
                "mean": mean_ct,
                "max": max_ct,
                "std": std_ct,
                "outliers": outlier_count
            }


        program_groups = all_cycles_df.groupby("program_name")

       
        analyses = {}
        for program_name, job_df in program_groups:
            analyses[program_name] = analyze_job(job_df.copy())

        results = []
        for program_name, analysis in analyses.items():

            results.append({
                "Program": program_name,
                "Cycles": analysis["count"],
                "Mean": analysis["mean"],
                "Median": analysis["median"],
                "Min": analysis["min"],
                "Max": analysis["max"],
                "STD": analysis["std"],
                "Outliers": analysis["outliers"],
                "Status": analysis["status"]
            })

        summary_df = pd.DataFrame(results)

        # =========================
        # Write to Excel 
        # =========================
        with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
            summary_df.to_excel(writer, sheet_name="Summary", index=False)
            for program_name, analysis, in analyses.items():
                sheet_name = str(program_name)[:31]
                analysis["df"].to_excel(writer, sheet_name=sheet_name, index=False)


        # =========================
        # Apply color to Highlight outliers 
        # =========================
        wb = load_workbook(OUTPUT_XLSX)

        red_fill = PatternFill(
            start_color="FFFF9999",
            end_color="FFFF9999",
            fill_type="solid"
        )

        for ws in wb.worksheets:

            if ws.title == "Summary":
                continue

            headers = [cell.value for cell in ws[1]]

            if "Outlier" not in headers:
                continue

            outlier_col = headers.index("Outlier") + 1

            for row in range(2, ws.max_row + 1):
                if ws.cell(row=row, column=outlier_col).value:
                    for col in range(1, ws.max_column + 1):

                        ws.cell(
                            row=row,
                            column=col
                        ).fill = red_fill

        wb.save(OUTPUT_XLSX)

        # =========================
        #Result window
        # =========================
        messagebox.showinfo(f"Success", #Title
                            f"Report Saved To:\n{OUTPUT_XLSX}") #Message

    # =========================
    # Exception if error
    # =========================
    except Exception as e:
        messagebox.showerror("Error", str(e))



# ===============================================
# GUI FUNCTIONS
# ===============================================

#Create close function
def close_app():
    root.destroy()

#Create user-selection window
def select_user():
    #Create window
    win = tk.Toplevel(root)
    win.title("Select User")
    win.geometry("300x150")

    #Prompt user to choose name
    tk.Label(win, text="Select User:").pack(pady=10)

    user_var = tk.StringVar()

    #Dropdown with usernames populated
    user_combo = ttk.Combobox(win, textvariable=user_var, values=list(users.keys()), state="readonly", width=25)
    user_combo.pack(pady=10)

    def save_user():
        #Allow updates to global variables
        global current_user
        global current_api_key

        #Get selected username from dropdown
        current_user = user_var.get()

        #Confirm user was selected
        if not current_user:
            messagebox.showwarning("No Selection", "Please select a user.")
            return

        #Get API key for the selected user
        current_api_key = users[current_user]

        #Refresh locations based on selected user
        refresh_locations()

        footer_label.config(text=f"Logged in as: {current_user}")

        #Open JSON file in "w"rite mode
        #with open("config.json", "w") as f:
        #    json.dump({"last_user": current_user}, f) #Convert Python into JSON file f

        save_last_user(current_user)

        print(f"Selected User: {current_user}")
        print(f"API Key: {current_api_key}")

        #Close window
        win.destroy()

    #Saves the selection
    tk.Button(win, text="Save", command=save_user).pack(pady=10)



# ===============================================
# GUI SETUP
# ===============================================

#Create main window
root = tk.Tk()
root.title("Drive Health Analyzer")
root.geometry("400x350")
root.configure(bg="#f0f0f0")


#Create Menu Bar
menu_bar = tk.Menu(root)

#Create Individual Menus
file_menu = tk.Menu(menu_bar, tearoff=0)
settings_menu = tk.Menu(menu_bar, tearoff=0)

#Attach menus to menu bar
menu_bar.add_cascade(label="File", menu=file_menu)
menu_bar.add_cascade(label="Settings", menu=settings_menu)

#Add commands
file_menu.add_command(label="Close", command=close_app)
settings_menu.add_command(label="Select User", command=select_user)

root.config(menu=menu_bar)


# Dropdown
tk.Label(root, text="Company", font=("Calibri",12)).pack()
company_var = tk.StringVar()
company_combo = ttk.Combobox(root, textvariable=company_var, values=[], state="readonly")
company_combo.pack(pady=5) 

if current_api_key:
    refresh_locations()

#Footer
footer_label = tk.Label(root, text = f"Logged in as: {current_user}", font = ("Calibri", 8))
footer_label.pack(side="bottom", pady=5)

#F-Number
tk.Label(root, text="F-Number", font=("Calibri",12)).pack()
serial_entry = tk.Entry(root)
serial_entry.pack()


#Dates
tk.Label(root, text="Start Date (YYYY-MM-DD)", font=("Calibri",12)).pack()
start_entry = tk.Entry(root)
start_entry.pack()

tk.Label(root, text="End Date (YYYY-MM-DD)", font=("Calibri",12)).pack()
end_entry = tk.Entry(root)
end_entry.pack()

# Run button
tk.Button(root, text="Run Report", command=run_report).pack(pady=20)

root.mainloop()
